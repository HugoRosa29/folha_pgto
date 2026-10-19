"""
Conferência da remuneração publicada no Portal da Transparência
contra a folha de pagamento do BENNER.

De um lado, os arquivos REMUNERACAO_<EMPRESA>_<AAAAMMDD>.txt
enviados ao portal (um por empresa, separados por ponto e
vírgula). Do outro, a planilha extraída do BENNER com a folha
das três empresas (ver "SELECT para emissão da Folha.docx").

A conferência casa os lançamentos por empresa + matrícula +
rubrica e separa:

    Valor diferente           o mesmo lançamento, por outro valor
    Tipo diferente            provento de um lado, desconto do outro
    Ausente do TXT            está na Folha e não foi publicado
    Ausente da Folha          foi publicado e não está na Folha
    Quantidade diferente      mesmo total, número de linhas diferente

Rubricas que nunca aparecem em nenhum TXT não são comparadas
linha a linha: elas ficam fora do portal por definição (planos
de saúde, empréstimos, bolsa de estágio...). Saem num resumo à
parte, com os proventos destacados, porque um provento fora do
portal merece uma segunda olhada.

Python puro: depende só do openpyxl, para rodar também no
navegador via Pyodide.
"""

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import unicodedata


CENTAVO = Decimal("0.01")

PROVENTO = 1
DESCONTO = 2

TIPO_VERBA = {PROVENTO: "Provento", DESCONTO: "Desconto"}

# Nome curto de cada empresa, pelo CNPJ que o TXT traz.
EMPRESAS = {
    "00359877000173": "TERRACAP",
    "50698494000106": "ETR",
    "29580134000100": "BIOTIC",
}

# Layout do TXT do portal: 30 campos separados por ";" (e um
# ";" final). É a mesma ordem do cabeçalho usado na conferência
# manual (Conf Portal Transparência).
CAMPOS_TXT = (
    "ID",
    "PAGADOR",
    "CD_EMPRESA",
    "DC_EMPRESA",
    "DC_GRUPO_SITFUNC",
    "MATRICULA",
    "NR_CPF",
    "CPF_FORMATADO",
    "NOME",
    "DC_CATEGORIA",
    "DC_CARGO",
    "DC_LOTACAO",
    "DC_FUNCAO",
    "SIT_OCUP",
    "ADMISSAO",
    "REG_JURIDICO",
    "CD_RUBRICA",
    "DC_RUBRICA",
    "NR_SEQUENCIA",
    "VL_RECEBIMENTO",
    "VL_DESCONTO",
    "SINDIRETA",
    "SINDIFICO",
    "CD_CARREIRA",
    "DC_CARREIRA",
    "CARGA_PRINCIPAL",
    "CARGA_SECUNDARIA",
    "CO_EXERCICIO",
    "CO_MES_REFERENCIA",
    "DS_NATUREZA",
)

# Colunas da planilha da Folha (saída do SELECT do BENNER).
COLUNAS_FOLHA = {
    "empresa": "EMPRESA",
    "grupo": "GRUPOPROCESSAMENTO",
    "tipo_colaborador": "TIPOCOLABROADOR",
    "matricula": "MATRICULA",
    "nome": "NOMEFUNCIONARIO",
    "competencia": "COMPETENCIA",
    "folha": "FOLHA",
    "codigo": "CODIGO",
    "rubrica": "NOME",
    "valor": "VALOR",
    "tipo_verba": "TIPOVERBA",
}

# Situações da conferência linha a linha, na ordem em que saem.
SITUACOES = (
    "valor_diferente",
    "tipo_diferente",
    "ausente_txt",
    "ausente_folha",
    "quantidade_diferente",
)

SITUACAO_TEXTO = {
    "conferido": "Conferido",
    "valor_diferente": "Valor diferente",
    "tipo_diferente": "Provento x desconto",
    "ausente_txt": "Na Folha, ausente do TXT",
    "ausente_folha": "No TXT, ausente da Folha",
    "quantidade_diferente": "Quantidade de lançamentos diferente",
}


# ============================================================
# UTILIDADES
# ============================================================

def normalizar(texto):
    """
    Maiúsculas, sem acento e com espaços simples: a forma usada
    para comparar nomes de empresa e de pessoa.
    """

    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))

    return " ".join(texto.upper().split())


def formatar_moeda(valor):
    """
    R$ 1.234,56 — o formato que se lê no relatório.
    """

    texto = f"{Decimal(valor or 0):,.2f}"

    return "R$ " + texto.replace(",", "_").replace(".", ",").replace("_", ".")


def _decimal(valor):
    """
    Valor monetário com duas casas. Aceita texto com ponto
    decimal (TXT) e número do Excel.
    """

    if valor is None or valor == "":
        return None

    if isinstance(valor, float):
        valor = repr(valor)

    try:
        return Decimal(str(valor).strip()).quantize(CENTAVO)
    except InvalidOperation:
        return None


def _matricula(valor):
    """
    A matrícula como texto, sem zeros à esquerda nem ".0" do
    Excel, para casar TXT e planilha.
    """

    if valor is None:
        return ""

    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)

    texto = str(valor).strip()

    return texto.lstrip("0") or texto


def _codigo(valor):
    """
    Código da rubrica como inteiro (ou texto, se não for número).
    """

    if valor is None:
        return None

    if isinstance(valor, float) and valor.is_integer():
        return int(valor)

    texto = str(valor).strip()

    return int(texto) if texto.isdigit() else texto


def _bytes(origem):
    """
    Aceita caminho, bytes ou objeto com .read() (arquivo do
    navegador ou upload).
    """

    if isinstance(origem, (bytes, bytearray)):
        return bytes(origem)

    if hasattr(origem, "read"):
        return origem.read()

    with open(origem, "rb") as arquivo:
        return arquivo.read()


def _nome_arquivo(origem):
    for atributo in ("filename", "name"):
        nome = getattr(origem, atributo, None)
        if nome:
            return str(nome).replace("\\", "/").rsplit("/", 1)[-1]

    if isinstance(origem, (str,)) or hasattr(origem, "__fspath__"):
        return str(origem).replace("\\", "/").rsplit("/", 1)[-1]

    return ""


def _texto(conteudo):
    """
    Os TXT do portal saem em ISO-8859-1; aceita UTF-8 também.
    """

    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue

    return conteudo.decode("latin-1", errors="replace")


# ============================================================
# LEITURA DO TXT DO PORTAL
# ============================================================

def ler_txt(origem, nome=None):
    """
    Lê um REMUNERACAO_<EMPRESA>_<AAAAMMDD>.txt.

    Devolve o arquivo com as linhas já interpretadas e a lista
    de problemas de leitura (linha fora do layout, valor
    ilegível...), que entram como inconsistências do próprio TXT.
    """

    nome = nome or _nome_arquivo(origem)
    texto = _texto(_bytes(origem))

    linhas = []
    problemas = []
    empresas = Counter()
    competencias = Counter()

    for numero, bruta in enumerate(texto.splitlines(), start=1):

        if not bruta.strip():
            continue

        campos = bruta.split(";")

        # O layout termina em ";", o que dá um campo vazio a mais.
        if len(campos) == len(CAMPOS_TXT) + 1 and campos[-1] == "":
            campos = campos[:-1]

        if len(campos) != len(CAMPOS_TXT):
            problemas.append({
                "arquivo": nome,
                "linha": numero,
                "problema": (
                    f"Linha com {len(campos)} campos; o layout tem "
                    f"{len(CAMPOS_TXT)}"
                ),
                "conteudo": bruta[:120],
            })
            continue

        registro = dict(zip(CAMPOS_TXT, (c.strip() for c in campos)))

        # Cabeçalho, quando o arquivo vier com um.
        if registro["CD_RUBRICA"].upper() == "CD_RUBRICA":
            continue

        recebimento = _decimal(registro["VL_RECEBIMENTO"])
        desconto = _decimal(registro["VL_DESCONTO"])

        if recebimento is not None and desconto is not None:
            problemas.append({
                "arquivo": nome,
                "linha": numero,
                "problema": "Linha com recebimento e desconto ao mesmo tempo",
                "conteudo": bruta[:120],
            })

        if recebimento is None and desconto is None:
            problemas.append({
                "arquivo": nome,
                "linha": numero,
                "problema": "Linha sem valor de recebimento nem de desconto",
                "conteudo": bruta[:120],
            })
            continue

        tipo = PROVENTO if recebimento is not None else DESCONTO
        valor = recebimento if recebimento is not None else desconto

        if valor <= 0:
            problemas.append({
                "arquivo": nome,
                "linha": numero,
                "problema": f"Valor zerado ou negativo ({valor})",
                "conteudo": bruta[:120],
            })

        cnpj = "".join(c for c in registro["CD_EMPRESA"] if c.isdigit()).zfill(14)
        empresa = EMPRESAS.get(cnpj) or registro["DC_EMPRESA"]

        empresas[(cnpj, registro["DC_EMPRESA"], empresa)] += 1
        competencias[(registro["CO_EXERCICIO"], registro["CO_MES_REFERENCIA"])] += 1

        linhas.append({
            "arquivo": nome,
            "linha": numero,
            "empresa": empresa,
            "empresa_nome": registro["DC_EMPRESA"],
            "cnpj": cnpj,
            "matricula": _matricula(registro["MATRICULA"]),
            "cpf": registro["NR_CPF"],
            "cpf_formatado": registro["CPF_FORMATADO"],
            "nome": registro["NOME"],
            "grupo": registro["DC_GRUPO_SITFUNC"],
            "cargo": registro["DC_CARGO"],
            "codigo": _codigo(registro["CD_RUBRICA"]),
            "rubrica": registro["DC_RUBRICA"],
            "natureza": registro["DS_NATUREZA"],
            "tipo": tipo,
            "valor": valor,
            "exercicio": registro["CO_EXERCICIO"],
            "mes": registro["CO_MES_REFERENCIA"],
        })

    if len(empresas) > 1:
        problemas.append({
            "arquivo": nome,
            "linha": None,
            "problema": (
                "O arquivo traz mais de uma empresa: "
                + ", ".join(sorted(e[2] for e in empresas))
            ),
            "conteudo": "",
        })

    principal = empresas.most_common(1)[0][0] if empresas else ("", "", "")

    return {
        "arquivo": nome,
        "empresa": principal[2],
        "empresa_nome": principal[1],
        "cnpj": principal[0],
        "competencias": sorted(competencias),
        "linhas": linhas,
        "problemas": problemas,
        "quantidade": len(linhas),
        "total_proventos": sum(
            (l["valor"] for l in linhas if l["tipo"] == PROVENTO), Decimal(0)
        ),
        "total_descontos": sum(
            (l["valor"] for l in linhas if l["tipo"] == DESCONTO), Decimal(0)
        ),
        "matriculas": len({l["matricula"] for l in linhas}),
    }


# ============================================================
# LEITURA DA FOLHA (BENNER)
# ============================================================

def ler_folha(origem, nome=None):
    """
    Lê a planilha da Folha das empresas, achando o cabeçalho
    pelos nomes das colunas (a ordem não importa).
    """

    from openpyxl import load_workbook

    nome = nome or _nome_arquivo(origem)
    livro = load_workbook(BytesIO(_bytes(origem)), read_only=True, data_only=True)

    try:
        for pagina in livro.worksheets:

            linhas = pagina.iter_rows(values_only=True)
            posicoes = None

            # O cabeçalho costuma ser a primeira linha; procura nas
            # dez primeiras por segurança.
            for _ in range(10):
                try:
                    cabecalho = next(linhas)
                except StopIteration:
                    break

                textos = [normalizar(c) for c in cabecalho]

                if all(coluna in textos for coluna in COLUNAS_FOLHA.values()):
                    posicoes = {
                        campo: textos.index(coluna)
                        for campo, coluna in COLUNAS_FOLHA.items()
                    }
                    break

            if posicoes is not None:
                return _ler_linhas_folha(nome, pagina.title, posicoes, linhas)
    finally:
        livro.close()

    raise ValueError(
        f"{nome}: não encontrei a tabela da Folha. As colunas esperadas "
        "são " + ", ".join(COLUNAS_FOLHA.values()) + "."
    )


def _ler_linhas_folha(nome, aba, posicoes, linhas):

    registros = []
    problemas = []
    competencias = Counter()
    folhas = Counter()

    for numero, linha in enumerate(linhas, start=2):

        if not linha or all(c is None or c == "" for c in linha):
            continue

        def campo(chave):
            indice = posicoes[chave]
            return linha[indice] if indice < len(linha) else None

        valor = _decimal(campo("valor"))
        tipo = _codigo(campo("tipo_verba"))

        if valor is None:
            problemas.append({
                "arquivo": nome,
                "linha": numero,
                "problema": f"Valor ilegível: {campo('valor')!r}",
                "conteudo": "",
            })
            continue

        competencia = campo("competencia")

        if isinstance(competencia, (datetime, date)):
            competencia = (str(competencia.year), str(competencia.month))
        else:
            competencia = (str(competencia or ""), "")

        competencias[competencia] += 1
        folhas[str(campo("folha") or "")] += 1

        registros.append({
            "arquivo": nome,
            "linha": numero,
            "empresa_nome": str(campo("empresa") or "").strip(),
            "grupo": str(campo("grupo") or "").strip(),
            "tipo_colaborador": str(campo("tipo_colaborador") or "").strip(),
            "matricula": _matricula(campo("matricula")),
            "nome": str(campo("nome") or "").strip(),
            "folha": str(campo("folha") or "").strip(),
            "codigo": _codigo(campo("codigo")),
            "rubrica": str(campo("rubrica") or "").strip(),
            "tipo": tipo,
            "valor": valor,
            "exercicio": competencia[0],
            "mes": competencia[1],
        })

    return {
        "arquivo": nome,
        "aba": aba,
        "linhas": registros,
        "problemas": problemas,
        "quantidade": len(registros),
        "competencias": sorted(competencias),
        "folhas": dict(folhas),
    }


# ============================================================
# CONFERÊNCIA
# ============================================================

def _empresa_da_folha(txts, folha):
    """
    Associa o nome da empresa na Folha ao nome curto do TXT,
    pelo nome por extenso (sem acento, maiúsculas).
    """

    por_nome = {normalizar(t["empresa_nome"]): t["empresa"] for t in txts}

    for cnpj, curto in EMPRESAS.items():
        por_nome.setdefault(normalizar(curto), curto)

    mapa = {}

    for linha in folha["linhas"]:
        chave = normalizar(linha["empresa_nome"])

        if chave not in mapa:
            mapa[chave] = por_nome.get(chave, linha["empresa_nome"])

        linha["empresa"] = mapa[chave]

    return mapa


def _resumo_lado(linhas):
    """
    Soma de um grupo de lançamentos com a mesma chave.
    """

    return {
        "quantidade": len(linhas),
        "valor": sum((l["valor"] for l in linhas), Decimal(0)),
        "tipos": sorted({l["tipo"] for l in linhas}),
        "valores": sorted(l["valor"] for l in linhas),
    }


def conferir(txts, folha):
    """
    Confere os TXT do portal com a Folha. Devolve um dicionário
    com cada bloco da conferência.
    """

    avisos = []
    mapa_empresas = _empresa_da_folha(txts, folha)

    empresas_txt = [t["empresa"] for t in txts]
    repetidas = [e for e, n in Counter(empresas_txt).items() if n > 1]

    for empresa in repetidas:
        avisos.append(
            f"Mais de um TXT da empresa {empresa}: os lançamentos foram "
            "somados. Confira se não é o mesmo arquivo enviado duas vezes."
        )

    empresas_folha = sorted(set(mapa_empresas.values()))
    sem_txt = [e for e in empresas_folha if e not in empresas_txt]
    sem_folha = [e for e in empresas_txt if e not in empresas_folha]

    for empresa in sem_txt:
        avisos.append(
            f"A Folha traz a empresa {empresa}, mas nenhum TXT dela foi "
            "enviado: os lançamentos dessa empresa não foram conferidos."
        )

    for empresa in sem_folha:
        avisos.append(
            f"O TXT da empresa {empresa} não tem nenhum lançamento na "
            "Folha enviada."
        )

    # Competência: o TXT e a Folha têm de ser do mesmo mês.
    comp_txt = sorted({c for t in txts for c in t["competencias"]})
    comp_folha = folha["competencias"]

    if comp_txt and comp_folha and set(comp_txt) != set(comp_folha):
        avisos.append(
            "Competências diferentes: TXT "
            + ", ".join(f"{m}/{a}" for a, m in comp_txt)
            + "; Folha "
            + ", ".join(f"{m}/{a}" for a, m in comp_folha)
            + "."
        )

    if len(folha["folhas"]) == 1:
        tipo_folha = next(iter(folha["folhas"]))
    else:
        tipo_folha = ", ".join(sorted(folha["folhas"]))

    conferidas = [e for e in empresas_folha if e in empresas_txt]

    linhas_txt = [l for t in txts for l in t["linhas"]]
    linhas_folha = [l for l in folha["linhas"] if l["empresa"] in conferidas]

    # Rubricas publicadas: as que aparecem em algum TXT. As demais
    # ficam fora do portal e não são comparadas linha a linha.
    publicadas = {l["codigo"] for l in linhas_txt}

    por_chave_txt = defaultdict(list)
    por_chave_folha = defaultdict(list)

    for linha in linhas_txt:
        por_chave_txt[(linha["empresa"], linha["matricula"], linha["codigo"])].append(linha)

    nao_publicadas = defaultdict(list)

    for linha in linhas_folha:
        if linha["codigo"] in publicadas:
            por_chave_folha[(linha["empresa"], linha["matricula"], linha["codigo"])].append(linha)
        else:
            nao_publicadas[linha["codigo"]].append(linha)

    # Quem é quem: nome e dados de cada colaborador, de cada lado.
    pessoas_txt = {}
    for linha in linhas_txt:
        pessoas_txt.setdefault((linha["empresa"], linha["matricula"]), linha)

    pessoas_folha = {}
    for linha in linhas_folha:
        pessoas_folha.setdefault((linha["empresa"], linha["matricula"]), linha)

    nomes_rubrica = {}
    for linha in linhas_folha + linhas_txt:
        nomes_rubrica.setdefault(linha["codigo"], linha["rubrica"])

    # ---------- linha a linha ----------

    divergencias = []
    conferidos = Counter()
    valor_conferido = Decimal(0)

    for chave in sorted(
        set(por_chave_txt) | set(por_chave_folha),
        key=lambda k: (k[0], k[1].zfill(12), str(k[2]).zfill(8)),
    ):
        empresa, matricula, codigo = chave
        lado_txt = por_chave_txt.get(chave, [])
        lado_folha = por_chave_folha.get(chave, [])

        # Colaborador inteiro fora de um dos lados vai para o bloco
        # de colaboradores; aqui ficam os lançamentos.
        pessoa = (lado_txt or lado_folha)[0]

        txt = _resumo_lado(lado_txt) if lado_txt else None
        fol = _resumo_lado(lado_folha) if lado_folha else None

        if txt and fol:
            if txt["tipos"] != fol["tipos"]:
                situacao = "tipo_diferente"
            elif txt["valor"] != fol["valor"]:
                situacao = "valor_diferente"
            elif txt["valores"] != fol["valores"]:
                situacao = "quantidade_diferente"
            else:
                situacao = "conferido"
        elif txt:
            situacao = "ausente_folha"
        else:
            situacao = "ausente_txt"

        if situacao == "conferido":
            conferidos[empresa] += 1
            valor_conferido += txt["valor"]
            continue

        tipo = (fol or txt)["tipos"][0] if (fol or txt)["tipos"] else None

        divergencias.append({
            "situacao": situacao,
            "empresa": empresa,
            "matricula": matricula,
            "nome": pessoa["nome"],
            "codigo": codigo,
            "rubrica": nomes_rubrica.get(codigo, ""),
            "tipo": TIPO_VERBA.get(tipo, str(tipo)),
            "tipo_txt": ", ".join(TIPO_VERBA.get(t, str(t)) for t in txt["tipos"]) if txt else "",
            "tipo_folha": ", ".join(TIPO_VERBA.get(t, str(t)) for t in fol["tipos"]) if fol else "",
            "quantidade_txt": txt["quantidade"] if txt else 0,
            "quantidade_folha": fol["quantidade"] if fol else 0,
            "valor_txt": txt["valor"] if txt else None,
            "valor_folha": fol["valor"] if fol else None,
            "diferenca": (txt["valor"] if txt else Decimal(0)) - (fol["valor"] if fol else Decimal(0)),
            "natureza": lado_txt[0]["natureza"] if lado_txt else "",
            "colaborador_ausente": (
                (empresa, matricula) not in pessoas_folha if situacao == "ausente_folha"
                else (empresa, matricula) not in pessoas_txt if situacao == "ausente_txt"
                else False
            ),
            "observacao": "",
        })

    # ---------- colaboradores de um lado só ----------

    colaboradores = []

    folha_por_pessoa = defaultdict(list)
    for linha in linhas_folha:
        folha_por_pessoa[(linha["empresa"], linha["matricula"])].append(linha)

    txt_por_pessoa = defaultdict(list)
    for linha in linhas_txt:
        txt_por_pessoa[(linha["empresa"], linha["matricula"])].append(linha)

    for chave in sorted(set(txt_por_pessoa) - set(folha_por_pessoa)):
        if chave[0] not in conferidas:
            continue

        lancamentos = txt_por_pessoa[chave]
        rubricas = sorted({str(l["codigo"]) for l in lancamentos}, key=lambda c: c.zfill(8))

        colaboradores.append({
            "lado": "Só no TXT",
            "empresa": chave[0],
            "matricula": chave[1],
            "nome": lancamentos[0]["nome"],
            "grupo": lancamentos[0]["grupo"],
            "cargo": lancamentos[0]["cargo"],
            "lancamentos": len(lancamentos),
            "proventos": sum((l["valor"] for l in lancamentos if l["tipo"] == PROVENTO), Decimal(0)),
            "descontos": sum((l["valor"] for l in lancamentos if l["tipo"] == DESCONTO), Decimal(0)),
            "rubricas": ", ".join(
                f"{l['codigo']} {l['rubrica']}"
                for l in sorted(lancamentos, key=lambda l: str(l["codigo"]).zfill(8))
            ),
            "relevante": True,
            "observacao": (
                f"Sem lançamentos na Folha enviada ({tipo_folha}). "
                "Verifique se o pagamento saiu em outra folha "
                "(rescisão, complementar, férias)."
            ),
        })

    for chave in sorted(set(folha_por_pessoa) - set(txt_por_pessoa)):
        lancamentos = folha_por_pessoa[chave]
        so_nao_publicadas = all(l["codigo"] not in publicadas for l in lancamentos)

        colaboradores.append({
            "lado": "Só na Folha",
            "empresa": chave[0],
            "matricula": chave[1],
            "nome": lancamentos[0]["nome"],
            "grupo": lancamentos[0]["grupo"],
            "cargo": lancamentos[0]["tipo_colaborador"],
            "lancamentos": len(lancamentos),
            "proventos": sum((l["valor"] for l in lancamentos if l["tipo"] == PROVENTO), Decimal(0)),
            "descontos": sum((l["valor"] for l in lancamentos if l["tipo"] == DESCONTO), Decimal(0)),
            "rubricas": ", ".join(
                f"{l['codigo']} {l['rubrica']}"
                for l in sorted(lancamentos, key=lambda l: str(l["codigo"]).zfill(8))
            ),
            "relevante": not so_nao_publicadas,
            "observacao": (
                "Recebe só rubricas que nenhum TXT publica"
                if so_nao_publicadas
                else "Tem rubricas publicadas para outros colaboradores"
            ),
        })

    # Os lançamentos dos colaboradores de um lado só já estão no
    # bloco acima; nas divergências, ficam marcados para filtro.
    for item in divergencias:
        if item["colaborador_ausente"]:
            item["observacao"] = (
                "Colaborador sem nenhum lançamento na Folha"
                if item["situacao"] == "ausente_folha"
                else "Colaborador sem nenhum lançamento no TXT"
            )

    # ---------- nomes ----------

    nomes_diferentes = []

    for chave in sorted(set(pessoas_txt) & set(pessoas_folha)):
        nome_txt = pessoas_txt[chave]["nome"]
        nome_folha = pessoas_folha[chave]["nome"]

        if normalizar(nome_txt) != normalizar(nome_folha):
            nomes_diferentes.append({
                "empresa": chave[0],
                "matricula": chave[1],
                "nome_txt": nome_txt,
                "nome_folha": nome_folha,
            })

    # ---------- rubricas fora do portal ----------

    rubricas_fora = []

    for codigo, lancamentos in nao_publicadas.items():
        tipos = sorted({l["tipo"] for l in lancamentos})
        por_empresa = Counter(l["empresa"] for l in lancamentos)
        grupos = Counter(l["grupo"] for l in lancamentos)

        rubricas_fora.append({
            "codigo": codigo,
            "rubrica": lancamentos[0]["rubrica"],
            "tipo": ", ".join(TIPO_VERBA.get(t, str(t)) for t in tipos),
            "provento": PROVENTO in tipos,
            "lancamentos": len(lancamentos),
            "colaboradores": len({(l["empresa"], l["matricula"]) for l in lancamentos}),
            "valor": sum((l["valor"] for l in lancamentos), Decimal(0)),
            "empresas": ", ".join(f"{e} ({n})" for e, n in sorted(por_empresa.items())),
            "grupos": ", ".join(g for g, _ in grupos.most_common()),
        })

    rubricas_fora.sort(key=lambda r: (not r["provento"], -r["valor"]))

    lancamentos_fora = sorted(
        (
            {
                "empresa": l["empresa"],
                "matricula": l["matricula"],
                "nome": l["nome"],
                "grupo": l["grupo"],
                "codigo": l["codigo"],
                "rubrica": l["rubrica"],
                "tipo": TIPO_VERBA.get(l["tipo"], str(l["tipo"])),
                "valor": l["valor"],
            }
            for linhas in nao_publicadas.values()
            for l in linhas
        ),
        key=lambda r: (r["tipo"] != "Provento", r["empresa"], r["codigo"] if isinstance(r["codigo"], int) else 0, r["nome"]),
    )

    # ---------- checagens internas do TXT ----------

    internos = _checar_txt(txts, linhas_txt, por_chave_folha)

    # ---------- totais por empresa ----------

    totais = []

    for empresa in sorted(set(empresas_txt) | set(empresas_folha)):
        txt_emp = [l for l in linhas_txt if l["empresa"] == empresa]
        fol_emp = [l for l in folha["linhas"] if l.get("empresa") == empresa]
        fol_pub = [l for l in fol_emp if l["codigo"] in publicadas]

        def soma(linhas, tipo):
            return sum((l["valor"] for l in linhas if l["tipo"] == tipo), Decimal(0))

        divergentes = [d for d in divergencias if d["empresa"] == empresa]

        totais.append({
            "empresa": empresa,
            "arquivo_txt": ", ".join(t["arquivo"] for t in txts if t["empresa"] == empresa),
            "linhas_txt": len(txt_emp),
            "linhas_folha": len(fol_emp),
            "linhas_folha_publicaveis": len(fol_pub),
            "colaboradores_txt": len({l["matricula"] for l in txt_emp}),
            "colaboradores_folha": len({l["matricula"] for l in fol_emp}),
            "proventos_txt": soma(txt_emp, PROVENTO),
            "proventos_folha": soma(fol_pub, PROVENTO),
            "descontos_txt": soma(txt_emp, DESCONTO),
            "descontos_folha": soma(fol_pub, DESCONTO),
            "proventos_folha_total": soma(fol_emp, PROVENTO),
            "descontos_folha_total": soma(fol_emp, DESCONTO),
            "conferidos": conferidos[empresa],
            "divergencias": len(divergentes),
            "por_situacao": dict(Counter(d["situacao"] for d in divergentes)),
            "comparada": empresa in conferidas,
        })

    contagem = Counter(d["situacao"] for d in divergencias)

    return {
        "avisos": avisos,
        "txts": [
            {chave: valor for chave, valor in t.items() if chave != "linhas"}
            for t in txts
        ],
        "folha": {chave: valor for chave, valor in folha.items() if chave != "linhas"},
        "tipo_folha": tipo_folha,
        "totais": totais,
        "divergencias": divergencias,
        "por_situacao": {s: contagem.get(s, 0) for s in SITUACOES},
        "quantidade_conferidos": sum(conferidos.values()),
        "valor_conferido": valor_conferido,
        "colaboradores": colaboradores,
        "nomes_diferentes": nomes_diferentes,
        "rubricas_fora": rubricas_fora,
        "lancamentos_fora": lancamentos_fora,
        "internos": internos,
        "quantidade_colaboradores_relevantes": sum(1 for c in colaboradores if c["relevante"]),
        "confere": (
            not divergencias
            and not nomes_diferentes
            and not any(c["relevante"] for c in colaboradores)
            and not internos["problemas"]
            and all(r["explicada"] for r in internos["repeticoes"])
        ),
    }


def _checar_txt(txts, linhas_txt, por_chave_folha):
    """
    O que dá para conferir no TXT sem a Folha: layout, CPF,
    cadastro coerente e lançamentos repetidos.
    """

    problemas = [dict(p) for t in txts for p in t["problemas"]]

    # CPF mascarado tem de bater com o CPF completo.
    for linha in linhas_txt:
        cpf = "".join(c for c in linha["cpf"] if c.isdigit()).zfill(11)
        esperado = "***" + cpf[3:9] + "**"

        if linha["cpf_formatado"] and linha["cpf_formatado"] != esperado:
            problemas.append({
                "arquivo": linha["arquivo"],
                "linha": linha["linha"],
                "problema": (
                    f"CPF_FORMATADO {linha['cpf_formatado']} não corresponde "
                    f"ao CPF {linha['cpf']} (esperado {esperado})"
                ),
                "conteudo": f"{linha['matricula']} {linha['nome']}",
            })

    # Mesma matrícula com cadastro diferente entre as linhas.
    cadastro = defaultdict(lambda: defaultdict(set))

    for linha in linhas_txt:
        pessoa = cadastro[(linha["empresa"], linha["matricula"])]
        for campo in ("cpf", "nome", "grupo", "cargo"):
            pessoa[campo].add(linha[campo])

    for (empresa, matricula), campos in sorted(cadastro.items()):
        for campo, valores in campos.items():
            if len(valores) > 1:
                problemas.append({
                    "arquivo": empresa,
                    "linha": None,
                    "problema": (
                        f"Matrícula {matricula} com {campo.upper()} diferente "
                        f"entre as linhas: " + " | ".join(sorted(valores))
                    ),
                    "conteudo": "",
                })

    # A mesma rubrica com naturezas diferentes no portal.
    naturezas = defaultdict(set)

    for linha in linhas_txt:
        naturezas[linha["codigo"]].add(linha["natureza"])

    for codigo, valores in sorted(naturezas.items(), key=lambda i: str(i[0]).zfill(8)):
        if len(valores) > 1:
            problemas.append({
                "arquivo": "",
                "linha": None,
                "problema": (
                    f"Rubrica {codigo} classificada em naturezas diferentes: "
                    + " | ".join(sorted(valores))
                ),
                "conteudo": "",
            })

    # Repetições de nome + rubrica, como na conferência manual
    # (Conf Portal Transparência), já com a explicação de cada uma.
    por_nome = defaultdict(list)

    for linha in linhas_txt:
        por_nome[(linha["empresa"], normalizar(linha["nome"]), linha["codigo"])].append(linha)

    repeticoes = []

    for (empresa, _, codigo), linhas in sorted(por_nome.items()):
        if len(linhas) < 2:
            continue

        matriculas = sorted({l["matricula"] for l in linhas})

        if len(matriculas) > 1:
            explicacao = f"{len(matriculas)} vínculos (matrículas {', '.join(matriculas)})"
            explicada = True
        else:
            na_folha = por_chave_folha.get((empresa, matriculas[0], codigo), [])
            valores_txt = sorted(l["valor"] for l in linhas)
            valores_folha = sorted(l["valor"] for l in na_folha)

            if valores_txt == valores_folha:
                explicacao = (
                    f"A Folha também traz {len(na_folha)} lançamentos "
                    "com os mesmos valores"
                )
                explicada = True
            else:
                explicacao = (
                    f"Repetido no TXT; a Folha traz {len(na_folha)} "
                    "lançamento(s) com outros valores"
                )
                explicada = False

        repeticoes.append({
            "empresa": empresa,
            "nome": linhas[0]["nome"],
            "matriculas": ", ".join(matriculas),
            "codigo": codigo,
            "rubrica": linhas[0]["rubrica"],
            "quantidade": len(linhas),
            "valores": " + ".join(formatar_moeda(l["valor"]) for l in linhas),
            "explicacao": explicacao,
            "explicada": explicada,
        })

    return {
        "problemas": problemas,
        "repeticoes": repeticoes,
    }


# ============================================================
# ENTRADA ÚNICA
# ============================================================

def comparar(txts, folha):
    """
    Lê os arquivos e faz a conferência.

    txts: lista de caminhos, bytes ou arquivos (um por empresa)
    folha: caminho, bytes ou arquivo da planilha da Folha
    """

    if not txts:
        raise ValueError("Envie pelo menos um TXT do portal.")

    if folha is None:
        raise ValueError("Envie a planilha da Folha.")

    lidos = [ler_txt(origem) for origem in txts]
    planilha = ler_folha(folha)

    return conferir(lidos, planilha)
