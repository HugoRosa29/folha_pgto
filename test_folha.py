"""
Testes da conferência Portal x Folha.

Os testes sintéticos montam TXT e planilha em memória e rodam em
qualquer lugar (inclusive no CI). Os testes com os arquivos reais
só rodam quando a pasta é indicada, para que nenhum dado de folha
precise ser copiado para o repositório:

    $env:FOLHA_REF = "C:\\Users\\hugos\\Downloads\\FOLHA_PGTO\\ARQUIVOS_REF"
    python -m unittest -v test_folha
"""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import os
import unittest

from exportar import montar_planilha
from folha import comparar, ler_txt


CABECALHO_FOLHA = [
    "EMPRESA", "GRUPOPROCESSAMENTO", "TIPOCOLABROADOR", "MATRICULA",
    "NOMEFUNCIONARIO", "COMPETENCIA", "FOLHA", "CODIGO", "NOME",
    "QUANTIDADE", "UNIDADE", "VALOR", "TIPOVERBA",
]

TERRACAP = ("00359877000173", "COMPANHIA IMOBILIÁRIA DE BRASÍLIA TERRACAP")


def linha_txt(matricula, nome, codigo, rubrica, recebimento="", desconto="",
              cpf="70997080159", empresa=TERRACAP):
    mascara = "***" + cpf[3:9] + "**"
    campos = [
        "", "", empresa[0], empresa[1], "ATIVO PERMANENTE", str(matricula),
        cpf, mascara, nome, "Contador", "Contador", "NUREC", "",
        "COM VINCULO", "20100201", "CLT", str(codigo), rubrica, "",
        recebimento, desconto, "N", "N", "553", "ATIVO PERMANENTE", "40",
        "", "2026", "9", "REMUNERAÇÃO BÁSICA",
    ]
    return ";".join(campos) + ";"


def txt(*linhas):
    return ("\r\n".join(linhas) + "\r\n").encode("latin-1")


def folha(*linhas):
    import xlsxwriter

    buffer = BytesIO()
    livro = xlsxwriter.Workbook(buffer, {"in_memory": True})
    pagina = livro.add_worksheet("Registros_1")
    data = livro.add_format({"num_format": "dd/mm/yyyy"})

    pagina.write_row(0, 0, CABECALHO_FOLHA)

    for numero, (matricula, nome, codigo, rubrica, valor, tipo) in enumerate(linhas, start=1):
        pagina.write_row(numero, 0, [
            TERRACAP[1], "Efetivos", "Funcionario", matricula, nome,
        ])
        pagina.write_datetime(numero, 5, datetime(2026, 9, 1), data)
        pagina.write_row(numero, 6, ["Folha Normal", codigo, rubrica, 0, 0, valor, tipo])

    livro.close()

    return buffer.getvalue()


class Sinteticos(unittest.TestCase):

    def test_tudo_confere(self):
        resultado = comparar(
            [txt(
                linha_txt(1, "Ana", 1, "Salário", recebimento="1000.5"),
                linha_txt(1, "Ana", 1500, "INSS", desconto="100"),
            )],
            folha(
                (1, "Ana", 1, "Salário", 1000.50, 1),
                (1, "Ana", 1500, "INSS", 100, 2),
                (1, "Ana", 1745, "GEAP", 300, 2),
            ),
        )

        self.assertTrue(resultado["confere"])
        self.assertEqual(resultado["quantidade_conferidos"], 2)
        self.assertEqual([r["codigo"] for r in resultado["rubricas_fora"]], [1745])

    def test_valor_tipo_e_ausencias(self):
        resultado = comparar(
            [txt(
                linha_txt(1, "Ana", 1, "Salário", recebimento="1000"),
                linha_txt(1, "Ana", 1500, "INSS", recebimento="100"),
                linha_txt(1, "Ana", 1600, "IRRF", desconto="50"),
                linha_txt(2, "Bia", 1, "Salário", recebimento="900"),
                linha_txt(9, "Rescindido", 1, "Salário", recebimento="10"),
            )],
            folha(
                (1, "Ana", 1, "Salário", 1000.01, 1),
                (1, "Ana", 1500, "INSS", 100, 2),
                (2, "Bia", 1, "Salário", 900, 1),
                (2, "Bia", 1600, "IRRF", 30, 2),
                (3, "Caio", 1, "Salário", 800, 1),
            ),
        )

        por_situacao = {
            (d["situacao"], d["matricula"], d["codigo"]) for d in resultado["divergencias"]
        }

        self.assertIn(("valor_diferente", "1", 1), por_situacao)
        self.assertIn(("tipo_diferente", "1", 1500), por_situacao)
        self.assertIn(("ausente_folha", "1", 1600), por_situacao)
        self.assertIn(("ausente_txt", "2", 1600), por_situacao)
        self.assertIn(("ausente_folha", "9", 1), por_situacao)
        self.assertIn(("ausente_txt", "3", 1), por_situacao)

        lados = {(c["lado"], c["matricula"], c["relevante"]) for c in resultado["colaboradores"]}
        self.assertEqual(lados, {("Só no TXT", "9", True), ("Só na Folha", "3", True)})
        self.assertFalse(resultado["confere"])

    def test_repeticao_e_apontada_mesmo_com_causa_provavel(self):
        resultado = comparar(
            [txt(
                linha_txt(5, "Maria", 106, "ATS", recebimento="638.65"),
                linha_txt(5, "Maria", 106, "ATS", recebimento="3193.23"),
            )],
            folha(
                (5, "Maria", 106, "ATS", 3193.23, 1),
                (5, "Maria", 106, "ATS", 638.65, 1),
            ),
        )

        # Bate com a Folha, mas a checagem da Conf aponta a repetição.
        self.assertEqual(resultado["divergencias"], [])
        self.assertFalse(resultado["confere"])

        repeticao, = resultado["internos"]["repeticoes"]
        self.assertEqual(repeticao["chave_conf"], "Maria106")
        self.assertTrue(repeticao["explicada"])

    def test_repeticao_sem_causa(self):
        resultado = comparar(
            [txt(
                linha_txt(5, "Maria", 106, "ATS", recebimento="100"),
                linha_txt(5, "Maria", 106, "ATS", recebimento="100"),
            )],
            folha((5, "Maria", 106, "ATS", 100, 1)),
        )

        repeticao, = resultado["internos"]["repeticoes"]
        self.assertFalse(repeticao["explicada"])

    def test_mesmo_total_com_lancamentos_diferentes(self):
        resultado = comparar(
            [txt(
                linha_txt(5, "Maria", 106, "ATS", recebimento="100"),
                linha_txt(5, "Maria", 106, "ATS", recebimento="100"),
            )],
            folha((5, "Maria", 106, "ATS", 200, 1)),
        )

        divergencia, = resultado["divergencias"]
        self.assertEqual(divergencia["situacao"], "quantidade_diferente")

    def test_estagiario_so_com_rubricas_fora_do_portal_e_informativo(self):
        resultado = comparar(
            [txt(linha_txt(1, "Ana", 1, "Salário", recebimento="1000"))],
            folha(
                (1, "Ana", 1, "Salário", 1000, 1),
                (7, "Estagiária", 7, "Bolsa estudo", 1000, 1),
            ),
        )

        self.assertTrue(resultado["confere"])
        colaborador, = resultado["colaboradores"]
        self.assertFalse(colaborador["relevante"])
        self.assertTrue(resultado["rubricas_fora"][0]["provento"])

    def test_problemas_de_layout_e_cpf(self):
        lido = ler_txt(txt(
            linha_txt(1, "Ana", 1, "Salário", recebimento="10", desconto="10"),
            "campo;a;menos",
            linha_txt(2, "Bia", 1, "Salário"),
        ), nome="teste.txt")

        problemas = [p["problema"] for p in lido["problemas"]]

        self.assertEqual(len(lido["linhas"]), 1)
        self.assertTrue(any("ao mesmo tempo" in p for p in problemas))
        self.assertTrue(any("campos" in p for p in problemas))
        self.assertTrue(any("sem valor" in p for p in problemas))

        linha = linha_txt(1, "Ana", 1, "Salário", recebimento="10").replace("***970801**", "***000000**")
        resultado = comparar([txt(linha)], folha((1, "Ana", 1, "Salário", 10, 1)))
        self.assertTrue(any("CPF_FORMATADO" in p["problema"] for p in resultado["internos"]["problemas"]))

    def test_exporta_excel(self):
        resultado = comparar(
            [txt(linha_txt(1, "Ana", 1, "Salário", recebimento="1000"))],
            folha((1, "Ana", 1, "Salário", 999, 1)),
        )

        conteudo = montar_planilha(resultado)

        from openpyxl import load_workbook

        livro = load_workbook(BytesIO(conteudo), read_only=True)
        self.assertIn("Divergencias", livro.sheetnames)
        self.assertIn("Rubricas fora do portal", livro.sheetnames)


PASTA_REAL = os.environ.get("FOLHA_REF")


@unittest.skipUnless(PASTA_REAL, "defina FOLHA_REF com a pasta dos arquivos reais")
class ArquivosReaisSetembro2026(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        from main import descobrir

        txts, planilha = descobrir(PASTA_REAL)
        cls.resultado = comparar(txts, planilha)

    def test_conferidos(self):
        self.assertEqual(self.resultado["quantidade_conferidos"], 4587)
        self.assertEqual(self.resultado["valor_conferido"], Decimal("22379045.65"))

    def test_so_no_txt_sao_as_quatro_rescisoes(self):
        so_txt = sorted(
            (c["empresa"], c["matricula"])
            for c in self.resultado["colaboradores"] if c["lado"] == "Só no TXT"
        )

        self.assertEqual(so_txt, [
            ("BIOTIC", "20000042"),
            ("TERRACAP", "10000072"),
            ("TERRACAP", "29319"),
            ("TERRACAP", "29831"),
        ])

        self.assertEqual(self.resultado["por_situacao"]["ausente_folha"], 31)
        self.assertTrue(all(d["colaborador_ausente"] for d in self.resultado["divergencias"]))

    def test_etr_fecha(self):
        etr, = [t for t in self.resultado["totais"] if t["empresa"] == "ETR"]
        self.assertEqual(etr["divergencias"], 0)
        self.assertEqual(etr["proventos_txt"], etr["proventos_folha"])

    def test_aponta_o_mesmo_que_a_conferencia_manual(self):
        terracap = {
            r["chave_conf"]: r["quantidade"]
            for r in self.resultado["internos"]["repeticoes"] if r["empresa"] == "TERRACAP"
        }

        # As 9 chaves "nome e verba para conferência" com QTDE 2 na
        # Planilha3 da Conf Portal Transparência TERRACAP.
        self.assertEqual(terracap, {
            "Adailton Silva Da Costa106": 2,
            "Adailton Silva Da Costa1": 2,
            "Jader Cabral De Almeida1600": 2,
            "Julio Cesar de Azevedo Reis1600": 2,
            "Maria Sitonio De Jesus Velasco5015": 2,
            "Maria Sitonio De Jesus Velasco5012": 2,
            "Maria Sitonio De Jesus Velasco106": 2,
            "Maria Sitonio De Jesus Velasco102": 2,
            "Maria Sitonio De Jesus Velasco10501": 2,
        })

        # ETR e BIOTIC não estavam na Conf; a mesma checagem acha mais duas.
        self.assertEqual(len(self.resultado["internos"]["repeticoes"]), 11)
        self.assertFalse(self.resultado["confere"])

    def test_nada_alem_disso(self):
        self.assertEqual(self.resultado["nomes_diferentes"], [])
        self.assertEqual(self.resultado["internos"]["problemas"], [])
        self.assertEqual(self.resultado["quantidade_colaboradores_relevantes"], 4)


if __name__ == "__main__":
    unittest.main()
